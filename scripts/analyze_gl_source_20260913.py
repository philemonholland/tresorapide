"""Read-only structural and cell audit for the supplied Grand Livre workbook."""

import json
import sys
from datetime import date, datetime, time
from decimal import Decimal

from openpyxl import load_workbook


def serializable(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value

def main(path, patterns=()):
    formulas = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    result = {
        "sheet_names": formulas.sheetnames,
        "defined_names": [name.name for name in formulas.defined_names.values()],
        "sheets": [],
    }
    for formula_sheet in formulas.worksheets:
        value_sheet = values[formula_sheet.title]
        selected_rows = set(range(1, min(formula_sheet.max_row, 30) + 1))
        if patterns:
            lowered_patterns = [pattern.casefold() for pattern in patterns]
            for source_row in formula_sheet.iter_rows():
                row_text = " | ".join(
                    str(cell.value or "") for cell in source_row
                ).casefold()
                if any(pattern in row_text for pattern in lowered_patterns):
                    selected_rows.update(
                        range(max(1, source_row[0].row - 3), min(formula_sheet.max_row, source_row[0].row + 3) + 1)
                    )
        else:
            selected_rows.update(range(1, formula_sheet.max_row + 1))
        rows = []
        for row_index in sorted(selected_rows):
            populated = []
            for column_index in range(1, formula_sheet.max_column + 1):
                formula_cell = formula_sheet.cell(row_index, column_index)
                value_cell = value_sheet.cell(row_index, column_index)
                if formula_cell.value is None and value_cell.value is None:
                    continue
                populated.append({
                    "coordinate": formula_cell.coordinate,
                    "formula_or_value": serializable(formula_cell.value),
                    "cached_value": serializable(value_cell.value),
                    "data_type": formula_cell.data_type,
                    "number_format": formula_cell.number_format,
                })
            if populated:
                rows.append({"row": row_index, "cells": populated})
        result["sheets"].append({
            "title": formula_sheet.title,
            "state": formula_sheet.sheet_state,
            "max_row": formula_sheet.max_row,
            "max_column": formula_sheet.max_column,
            "merged_ranges": [str(rng) for rng in formula_sheet.merged_cells.ranges],
            "rows": rows,
        })
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2:])
